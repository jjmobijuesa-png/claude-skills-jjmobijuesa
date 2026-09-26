# -*- coding: utf-8 -*-
"""Lee el chat de WhatsApp, extrae los reportes que empiezan con 'gracias totales',
parsea los renglones <monto> <concepto>, clasifica en rubros y registra en Excel.

SOLO LECTURA sobre WhatsApp. Nunca escribe ni envia mensajes.
Uso:  python leer_gastos_wa.py "<texto que identifica el chat>" [--scroll N]
"""
import sys, os, time, json, re, unicodedata, datetime
sys.stdout.reconfigure(encoding="utf-8")
from playwright.sync_api import sync_playwright

CDP="http://127.0.0.1:9222"
BASE=r"E:\vars\var 171 9 Finanza Integral"
DEST=os.path.join(BASE,"02 Egresos","Registro WhatsApp")
XLSX=os.path.join(DEST,"registro_gastos_wa.xlsx")
MARCADOR="gracias totales"

EXTRACT="""()=>{const out=[];
  document.querySelectorAll('#main [data-pre-plain-text]').forEach(e=>{
    const m=e.getAttribute('data-pre-plain-text')||'';
    const t=(e.innerText||'').trim();
    if(t) out.push({meta:m,txt:t});});
  return out;}"""
RX_META=re.compile(r'\[(\d{1,2}):(\d{2})\s*([ap])\.?\s*m\.?,\s*(\d{1,2})/(\d{1,2})/(\d{4})\]\s*(.*?):', re.I)
# monto: primer numero o cadena sumada (3.75 + 3.75)
RX_MONTO=re.compile(r'^\s*(\d+(?:[.,]\d+)?(?:\s*\+\s*\d+(?:[.,]\d+)?)*)\s+(.*)$')

def norm(s):
    s=unicodedata.normalize('NFKD',s or '')
    return ''.join(c for c in s if not unicodedata.combining(c)).lower().strip()

# --- clasificacion por rubro ---
FIJO=[('mutualista','Cuota Mutualista Pichincha'),('iess','IESS'),('peruzzi','Cuota Peruzzi'),
 ('cpa','Honorarios contables CPA'),('contador','Honorarios contables CPA'),
 ('cnel','Servicios basicos - CNEL'),('cnt','Servicios basicos - CNT'),('luz','Servicios basicos - CNEL'),
 ('telefon','Servicios basicos - CNT'),('internet','Servicios basicos - CNT'),
 ('botellon','Botellones de agua'),('agua','Servicios basicos - agua'),
 ('devolucion','Pago por devolucion'),('cuota','Cuota de credito'),('prestamo','Cuota de credito'),
 ('la nuestra','Cuota COAC La Nuestra'),('mobijuesa','Cuota anticipo Mobijuesa')]
VARIABLE=[('restaurante','Comida fuera de casa'),('comida fuera','Comida fuera de casa'),
 ('comer fuera','Comida fuera de casa'),('cine','Recreacion'),('vehiculo','Arreglo de vehiculo'),
 ('carro','Arreglo de vehiculo'),('auto','Arreglo de vehiculo'),('llanta','Arreglo de vehiculo'),
 ('ropa','Ropa'),('viaje','Viajes'),('peaje','Peajes'),('anthropic','Pago IA'),
 ('inteligencia artificial','Pago IA'),('copia','Copias'),('libro','Copias'),
 ('combustible','Combustible'),('gasolina','Combustible'),('diesel','Combustible'),
 ('mascota','Mascotas'),('medicin','Medicina'),('farmacia','Medicina'),('salud','Medicina'),
 ('alimento','Alimentos vivienda'),('leche','Alimentos vivienda'),('huevo','Alimentos vivienda'),
 ('queso','Alimentos vivienda'),('fruta','Alimentos vivienda'),('mercado','Alimentos vivienda'),
 ('mama','Aporte a Mama'),('parqueo','Parqueo'),('taxi','Peajes'),('donacion','Donacion'),
 ('corte','Corte de cabello'),('cabello','Corte de cabello'),('barber','Corte de cabello'),
 ('jardin','Arreglo Vivienda - jardin'),('planta','Arreglo Vivienda - jardin'),
 ('hogar','Utensilios de hogar'),('casa','Arreglo Vivienda'),('vivienda','Arreglo Vivienda'),
 ('limpieza','Utensilios de hogar'),('aseo','Utensilios de hogar'),('mascotas','Mascotas'),
 ('hijo','Hijo'),('gas','Alimentos vivienda'),('internet','Servicios basicos - CNT')]

# Inversion / proyecto (rubros historicos del flujo "Gracias 2026")
INVERSION=[('blockchain','Fundacion Blockchain'),('cripto','Criptomonedas'),('moringa','Moringa'),
 ('plano','Imp. Planos'),('patente','Patente'),('predio','Predios'),('alcabala','Alcabalas'),
 ('escritura','Escrituras'),('registro de la propiedad','Registro de la Propiedad'),
 ('permiso','Permiso Construccion'),('comision','Comisiones'),('importacion','Importacion'),
 ('limpieza de terreno','Limp. de Terreno y Cerco'),('cerco','Limp. de Terreno y Cerco'),
 ('demarcacion','Belen - demarcacion'),('shauly','Activacion canchas Shauly'),
 ('tasa','Tasas'),('notaria','Escrituras'),('municipio','Tasas'),
 ('inen','SAS Agua - normativa INEN'),('normativa','SAS Agua - normativa'),
 ('pozo','SAS Agua - pozo'),('surtidor','SAS Agua - equipo'),('dispensador','SAS Agua - equipo'),
 ('botellon','SAS Agua - insumos'),('filtro','SAS Agua - insumos')]

# PROYECTO: marcadores inequivocos de inversion en proyecto. Se evaluan ANTES que todo,
# porque de lo contrario una palabra generica ('agua') los captura como gasto corriente.
PROYECTO=[('inen','SAS Agua - normativa INEN'),('sas agua','SAS Agua - inversion'),
 ('punto de agua','SAS Agua - inversion'),('surtidor','SAS Agua - equipo'),
 ('dispensador','SAS Agua - equipo'),('purificadora','SAS Agua - equipo'),
 ('osmosis','SAS Agua - equipo'),('pozo','SAS Agua - pozo'),
 ('planta de agua','SAS Agua - equipo'),('normativa','SAS Agua - normativa')]

def clasifica(concepto):
    c=norm(concepto)
    for k,r in PROYECTO:
        if k in c: return 'INVERSION',r
    for k,r in FIJO:
        if k in c: return 'FIJO',r
    for k,r in INVERSION:
        if k in c: return 'INVERSION',r
    for k,r in VARIABLE:
        if k in c: return 'VARIABLE',r
    return 'RUBRO ABIERTO','Rubro abierto'

def parse_monto(s):
    partes=[float(x.replace(',','.')) for x in re.split(r'\s*\+\s*',s)]
    return round(sum(partes),2), partes

def main():
    chat = sys.argv[1] if len(sys.argv)>1 else "Mobijuesa"
    nsc  = 25
    if "--scroll" in sys.argv: nsc=int(sys.argv[sys.argv.index("--scroll")+1])
    os.makedirs(DEST, exist_ok=True)
    with sync_playwright() as p:
        try: br=p.chromium.connect_over_cdp(CDP)
        except Exception as e:
            print("No hay sesion CDP en 9222. Arranque WhatsApp con la skill "
                  "whatsapp-web-cdp-lectura-envio (wa_connect.py) y reintente."); return 2
        pgs=[x for x in br.contexts[0].pages if "web.whatsapp.com" in x.url]
        if not pgs: print("No hay pestana de WhatsApp abierta."); return 3
        # REGLA DE UNA SOLA PESTANA
        for extra in pgs[1:]:
            try: extra.close()
            except: pass
        pg=pgs[0]; pg.bring_to_front(); time.sleep(2)
        # esperar sincronizacion (puede tardar minutos)
        for k in range(40):
            if pg.query_selector('#pane-side'): break
            b=norm(pg.evaluate("()=>document.body.innerText") or '')
            if 'usar aqui' in b or 'use here' in b:
                print("'abierto en otra ventana': recargo SIN clicar 'Usar aqui'"); pg.reload(); time.sleep(6); continue
            if 'escanea' in b or 'scan the qr' in b: print("Pide QR: el usuario debe escanear una vez."); return 4
            if k%5==0: print(f"  esperando sincronizacion... {k*3}s")
            time.sleep(3)
        if not pg.query_selector('#pane-side'): print("No cargo la lista de chats."); return 5
        # localizar el chat recorriendo la lista (mas fiable que el buscador)
        filas=pg.query_selector_all('#pane-side [role="row"]'); idx=None
        for i,f in enumerate(filas):
            try:
                if norm(chat) in norm(f.inner_text() or ''): idx=i; break
            except: pass
        if idx is None: print(f"No encontre un chat que contenga '{chat}'."); return 6
        bx=filas[idx].bounding_box()
        pg.mouse.click(bx["x"]+bx["width"]/2, bx["y"]+bx["height"]/2); time.sleep(4)
        head=pg.evaluate("""()=>{const h=document.querySelector('#main header');return h?(h.innerText||'').trim():'';}""")
        print("CHAT ABIERTO:", head.replace("\n"," | ")[:100])
        # cargar historial
        acc={}
        pg.mouse.move(900,400)
        for k in range(nsc):
            for it in pg.evaluate(EXTRACT): acc[it['meta']+'|'+it['txt'][:80]]=it
            pg.mouse.wheel(0,-2200); time.sleep(1.1)
        for k in range(nsc+6): pg.mouse.wheel(0,4000); time.sleep(0.3)
        time.sleep(1.5)
        for it in pg.evaluate(EXTRACT): acc[it['meta']+'|'+it['txt'][:80]]=it
        print(f"Mensajes leidos: {len(acc)}")
    # --- filtrar por marcador y parsear ---
    filas_out=[]
    for it in acc.values():
        txt=it['txt']
        if not norm(txt).startswith(norm(MARCADOR)): continue
        m=RX_META.search(it['meta'])
        if m:
            hh=int(m.group(1)); mm=m.group(2)
            if m.group(3).lower()=='p' and hh!=12: hh+=12
            if m.group(3).lower()=='a' and hh==12: hh=0
            fecha=f"{int(m.group(6)):04d}-{int(m.group(5)):02d}-{int(m.group(4)):02d}"
            hora=f"{hh:02d}:{mm}"; autor=m.group(7).strip()
        else:
            fecha='?'; hora='?'; autor='?'
        lineas=[l.strip() for l in txt.splitlines()[1:] if l.strip()]
        for ln in lineas:
            if re.match(r'^\d{1,2}:\d{2}\s*[ap]', norm(ln)): continue   # marca de hora
            mo=RX_MONTO.match(ln)
            if not mo: continue
            monto,partes=parse_monto(mo.group(1))
            concepto=re.sub(r'\([^)]*\)','',mo.group(2)).strip(' .')
            tipo,rubro=clasifica(concepto)
            # CLAVE ESTABLE: se calcula del mensaje ORIGINAL, nunca de las celdas editables.
            # Asi, corregir un monto o un concepto en el Excel NO provoca reinsercion.
            import hashlib
            clave=hashlib.sha1(f"{fecha}|{hora}|{mo.group(1)}|{mo.group(2)}".encode('utf-8')).hexdigest()[:16]
            filas_out.append({'fecha':fecha,'hora':hora,'autor':autor,'monto':monto,
                              'desglose':(' + '.join(str(x) for x in partes) if len(partes)>1 else ''),
                              'concepto':concepto,'tipo':tipo,'rubro':rubro,'origen':'WhatsApp','clave':clave})
    if not filas_out:
        print(f"\nNo hay mensajes que empiecen con '{MARCADOR}'."); return 0
    filas_out.sort(key=lambda x:(x['fecha'],x['hora']))
    print(f"\n=== {len(filas_out)} RENGLONES DE GASTO ===")
    print(f"{'FECHA':11} {'HORA':6} {'MONTO':>9}  {'TIPO':14} {'RUBRO':28} CONCEPTO")
    tot=0
    for f in filas_out:
        print(f"{f['fecha']:11} {f['hora']:6} {f['monto']:9.2f}  {f['tipo']:14} {f['rubro'][:28]:28} {f['concepto'][:40]}")
        tot+=f['monto']
    print(f"{'':28}{tot:9.2f}  TOTAL")
    # --- registrar en Excel sin duplicar ---
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        if os.path.exists(XLSX):
            wb=openpyxl.load_workbook(XLSX); ws=wb.active
            # la clave estable vive en la columna J; si el libro es viejo y no la tiene, se crea
            if (ws.cell(1,10).value or '').strip().upper()!='CLAVE':
                c=ws.cell(1,10,'CLAVE'); c.font=Font(bold=True,color='FFFFFF')
                c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(horizontal='center')
            claves={str(ws.cell(r,10).value) for r in range(2,ws.max_row+1) if ws.cell(r,10).value}
        else:
            wb=openpyxl.Workbook(); ws=wb.active; ws.title='Gastos WA'
            for i,h in enumerate(['FECHA','HORA','AUTOR','MONTO','CONCEPTO','DESGLOSE','TIPO','RUBRO','ORIGEN','CLAVE'],1):
                c=ws.cell(1,i,h); c.font=Font(bold=True,color='FFFFFF')
                c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(horizontal='center')
            for col,w in zip('ABCDEFGHIJ',[12,8,18,11,38,14,15,28,11,18]): ws.column_dimensions[col].width=w
            ws.freeze_panes='A2'; claves=set()
        nuevos=0
        for f in filas_out:
            if f['clave'] in claves: continue
            r=ws.max_row+1
            for i,v in enumerate([f['fecha'],f['hora'],f['autor'],f['monto'],f['concepto'],
                                  f['desglose'],f['tipo'],f['rubro'],f['origen'],f['clave']],1):
                cc=ws.cell(r,i,v)
                if i==4: cc.number_format='#,##0.00'
            claves.add(f['clave']); nuevos+=1
        wb.save(XLSX)
        print(f"\nRegistro actualizado: {XLSX}")
        print(f"Renglones nuevos: {nuevos} | ya existentes (omitidos): {len(filas_out)-nuevos}")
    except Exception as e:
        print("No pude escribir el Excel:",e)
    return 0

if __name__=="__main__":
    sys.exit(main())
