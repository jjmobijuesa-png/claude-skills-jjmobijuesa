# -*- coding: utf-8 -*-
"""
Detector de MARCADORES DE SUBORDINACIÓN en un texto en español.

No juzga la cortesía (que es deseable): detecta la estructura profunda que
coloca al emisor por debajo del receptor —pedir permiso en vez de proponer
entre iguales— y propone la reescritura.

Uso:
    python revisar_sumision.py "ruta\\del\\borrador.txt|.md|.docx"
    python revisar_sumision.py --texto "Quisiera solicitarle que considere..."

Salida: hallazgos por categoría + índice de subordinación + veredicto.
La cortesía se conserva SIEMPRE; lo que se corrige es la posición.
"""
import sys, re
from pathlib import Path
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# (categoria, peso, patron, por_que, reescritura)
REGLAS = [
    ("modal-condicional", 3, r"\b(quisiera|querría|desearía|me gustaría|quisiéramos|desearíamos)\b",
     "verbo en condicional de deseo: pide permiso para querer",
     "usar presente de indicativo: «propongo», «planteamos»"),
    ("modal-condicional", 3, r"\b(podría[ns]?|pudiera[ns]?|sería posible|de ser posible|si fuera posible|si es posible)\b",
     "somete la acción a la posibilidad que concede el otro",
     "afirmar el hecho: «el arranque es en septiembre»"),
    ("ruego", 4, r"\b(agradecer[íi]a|agradecemos de antemano|le agradecer[íe]|mucho agradecer[ée])\b",
     "agradece por adelantado un favor: instala deuda antes de empezar",
     "agradecer al final y por algo ya ocurrido, o suprimir"),
    ("ruego", 4, r"\b(humildemente|respetuosamente solicito|acudo a usted|me permito solicitar|tenga a bien|sírvase)\b",
     "fórmula de súplica administrativa: posición de subordinado",
     "«proponemos», «presentamos», «acordamos»"),
    ("disculpa", 4, r"\b(disculpe (la|las) (molestia|interrupci[óo]n)|perd[óo]n por (la|las)|lamento molestar|s[ée] que est[áa] muy ocupad)\w*",
     "se disculpa por existir: cede el marco antes de hablar",
     "suprimir; ir al asunto"),
    ("permiso", 4, r"\b(solicit(o|amos|ar)\s+(su|la)\s+(autorizaci[óo]n|aprobaci[óo]n|permiso|venia|visto bueno)|pedir(le)?\s+(permiso|autorizaci[óo]n))\b",
     "convierte al otro en autoridad que concede",
     "«sometemos a su decisión las opciones A y B»"),
    ("espera-pasiva", 3, r"\b(quedo a la espera|quedamos a la espera|en espera de su|a la espera de sus|esperando su (pronta )?respuesta|quedo atento a lo que dispong)\w*",
     "cierra cediendo el control del tiempo",
     "cerrar con paso y fecha propios: «le confirmo el jueves»"),
    ("espera-pasiva", 2, r"\b(cuando (usted )?(lo )?(considere|estime|disponga|tenga a bien)|a su (entera )?disposici[óo]n|cuando guste)\b",
     "entrega la agenda al receptor",
     "proponer fecha concreta"),
    ("minimizacion", 3, r"\b(solo (quer[íi]a|quisiera)|apenas|un peque[ñn]o favor|una peque[ñn]a (consulta|duda)|si no es (mucha )?molestia|nada m[áa]s)\b",
     "minimiza el propio asunto: le baja valor antes de exponerlo",
     "nombrar el asunto por su tamaño real"),
    ("inseguridad", 3, r"\b(creo que tal vez|no s[ée] si|espero (que )?(no )?(le )?(sea|resulte)|ojal[áa]|si me permite|corr[íi]jame si)\b",
     "duda performativa: invita al receptor a dudar también",
     "afirmar y, si hay incertidumbre real, cuantificarla"),
    ("subordinacion-cargo", 2, r"\b(su distinguida|su honorable|tan digno|magn[íi]fic[oa]|ilustr[íi]sim)\w*",
     "asimetría honorífica: eleva al otro y baja al emisor",
     "tratamiento cortés simétrico: «Estimado Ing. X»"),
    ("futuro-condicional", 2, r"\b(estar[íi]amos (dispuestos|interesados)|nos gustar[íi]a poder|podr[íi]amos llegar a)\b",
     "condicional de disposición: no compromete, ruega",
     "«iniciamos», «entregamos», «operamos»"),
]

# marcadores POSITIVOS de posición entre iguales (suman confianza)
POSITIVOS = [
    (r"\b(propon(emos|go)|plante(amos|o)|present(amos|o)|acord(amos|é)|confirm(amos|o))\b", "verbo de acción propia en presente"),
    (r"\b(opci[óo]n (A|B|1|2)|dos (caminos|alternativas|opciones)|escenario (A|B))\b", "ofrece decisión entre opciones, no un sí/no"),
    (r"\b(el (\d{1,2}) de [a-záéíóú]+|antes del|a m[áa]s tardar|esta semana|el jueves|el lunes)\b", "fecha propia comprometida"),
    (r"\b(lo que (esto )?(les|le) (aporta|significa|resuelve)|a cambio|por su parte (reciben|obtienen))\b", "reciprocidad explícita en moneda del otro"),
]

def leer(ruta):
    p = Path(ruta)
    if p.suffix.lower() == ".docx":
        try:
            import docx
            d = docx.Document(str(p))
            t = [q.text for q in d.paragraphs]
            for tb in d.tables:
                for r in tb.rows:
                    t += [c.text for c in r.cells]
            return "\n".join(t)
        except Exception as e:
            print("  (no se pudo leer docx:", e, ")"); return ""
    return p.read_text(encoding="utf-8", errors="replace")

def analizar(texto):
    lineas = texto.splitlines()
    hallazgos = []
    for cat, peso, pat, porque, fix in REGLAS:
        for m in re.finditer(pat, texto, re.I):
            n = texto.count("\n", 0, m.start()) + 1
            frag = lineas[n-1].strip()[:110] if n-1 < len(lineas) else ""
            hallazgos.append((cat, peso, m.group(0), porque, fix, n, frag))

    pos = []
    for pat, desc in POSITIVOS:
        c = len(re.findall(pat, texto, re.I))
        if c: pos.append((desc, c))

    palabras = max(1, len(texto.split()))
    bruto = sum(h[1] for h in hallazgos)
    # indice normalizado por cada 300 palabras
    indice = round(bruto * 300 / palabras, 1)

    print("=" * 68)
    print("REVISION DE POSICION — ¿peticion de permiso o comunicacion entre iguales?")
    print("=" * 68)
    print(f"  Palabras: {palabras}  ·  marcadores de subordinacion: {len(hallazgos)}")
    print()
    if hallazgos:
        por_cat = {}
        for h in hallazgos:
            por_cat.setdefault(h[0], []).append(h)
        for cat, items in sorted(por_cat.items(), key=lambda x: -sum(i[1] for i in x[1])):
            print(f"  [{cat.upper()}]  {items[0][3]}")
            print(f"      -> {items[0][4]}")
            for it in items[:4]:
                print(f"         linea {it[5]}: «{it[2]}»   {it[6][:70]}")
            print()
    else:
        print("  Sin marcadores de subordinacion. \n")

    if pos:
        print("  MARCADORES DE POSICION ENTRE IGUALES (positivos):")
        for d, c in pos:
            print(f"      +{c}  {d}")
        print()

    if indice == 0:      v = "ENTRE IGUALES"
    elif indice <= 3:    v = "ACEPTABLE (revisar los marcados)"
    elif indice <= 8:    v = "PIDE PERMISO — reescribir"
    else:                v = "SUPLICA — reescribir por completo"
    print(f"  INDICE DE SUBORDINACION: {indice} por 300 palabras  ->  {v}")
    print()
    print("  Recordatorio: la CORTESIA se conserva siempre. Lo que se corrige es la POSICION.")
    print("  Y la posicion solo se sostiene si el contenido es verdadero y la peticion es explicita.")
    return indice, v

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    if sys.argv[1] == "--texto":
        analizar(" ".join(sys.argv[2:]))
    else:
        t = leer(sys.argv[1])
        if t: analizar(t)
