# -*- coding: utf-8 -*-
import re, glob, os, collections
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FLUJO=r"E:\vars\var 11-11 Finanza Integral\03 Flujo de Caja\Flujo de Efectivo 2026.xlsx"
DEST=r"E:\vars\var 11-11 Finanza Integral\10 Documentos de Respaldo\Estados de Cuenta\Produbanco 2026"
MES={'Jan':1,'Ene':1,'Feb':2,'Mar':3,'Apr':4,'Abr':4,'May':5,'Jun':6,'Jul':7,'Aug':8,'Ago':8,
     'Sep':9,'Set':9,'Oct':10,'Nov':11,'Dec':12,'Dic':12}
NOM={1:'Enero',2:'Febrero',3:'Marzo',4:'Abril',5:'Mayo',6:'Junio'}
LINE=re.compile(r'^([A-ZÁÉÍÓÚÑ]{3,})\s+([A-Za-z]{3})\s+(\d{1,2})\s+(\d{4})\s+(\S+)\s+(.+?)\s+([\d.,]+)([+-])\s+([\d.,]+)\s+([\d.,]+)\s*$')
num=lambda s: float(s.replace(',',''))
REGLAS=[
 ('TARIFA','Comisiones bancarias'),('REDONDEO','Aporte / Ahorro'),('AHORRO INCREMENTAL','Aporte / Ahorro'),
 ('FLEXIAHORRO','Aporte / Ahorro'),('AHORRAR POR AHORRAR','Aporte / Ahorro'),
 ('RETIROS ATM','** Retiro efectivo (revisar) **'),('RETIRO DE EFECTIVO','** Retiro efectivo (revisar) **'),
 ('INTERBANCARIA','** Transferencia (revisar) **'),('TRANSFERENCIA CTAS TERCEROS','** Transferencia (revisar) **'),
 ('TRANSFERENCIA ENTRE CTAS','** Transferencia (revisar) **'),('TRANSFERENCIA CTAS PROPIAS','** Transferencia (revisar) **'),
 ('CNEL','Servicios Basicos'),('AGUA POTABLE','Servicios Basicos'),('EMAPA','Servicios Basicos'),
 ('CNT','Internet CNT'),('CLARO','Movil Claro'),('MOVISTAR','Movil Claro'),('NETLIFE','Internet CNT'),
 ('PRIMAX','Combustible'),('PETROECUADOR','Combustible'),('MASGAS','Combustible'),('PRODUGAS','Combustible'),('GASOLINERA','Combustible'),
 ('ROSADO','Alimentos vivienda'),('SUPERMAXI','Alimentos vivienda'),('MEGAMAXI','Alimentos vivienda'),('COMISARIATO','Alimentos vivienda'),
 ('SANTA MARIA','Alimentos vivienda'),('GRAN AKI','Alimentos vivienda'),('TIA','Alimentos vivienda'),('CARNE','Alimentos vivienda'),('PANADERIA','Alimentos vivienda'),
 ('FARMACIA','Medicina'),('FYBECA','Medicina'),('SANA SANA','Medicina'),('PHARMACY','Medicina'),('CRUZ AZUL','Medicina'),
 ('MEDICITY','Medicina'),('DIFARE','Medicina'),('CLINICA','Medicina'),('HOSPITAL','Medicina'),('LABORATORIO','Medicina'),('OPTICA','Medicina'),
 ('SUPERDEPORTE','Ropa'),('DE PRATI','Ropa'),('ETAFASHION','Ropa'),('MARATHON','Ropa'),
 ('TECNICENTRO','Vehiculo (seguro-aceite)'),('LUBRICANT','Vehiculo (seguro-aceite)'),('LLANTA','Vehiculo (seguro-aceite)'),
 ('PARQUEO','Parqueo'),('PARKING','Parqueo'),('PEAJE','Peajes/Viajes'),('TAXI','Peajes/Viajes'),
 ('IESS','IESS'),('DIVIDENDO','Servicio de Deuda'),('PRESTAMO','Servicio de Deuda'),('AMORTIZA','Servicio de Deuda'),
 ('MUTUALISTA','Servicio de Deuda'),('LA NUESTRA','Servicio de Deuda'),('MOBIJUESA','Servicio de Deuda'),('COOPERATIVA','Servicio de Deuda'),
 ('NIPPONFLEX','Oficina - Nipponflex'),('DONACION','Donacion'),('DIEZMO','Donacion'),('IGLESIA','Donacion'),
 ('DEPOSITO','** INGRESO (revisar) **'),('TRANSFERENCIA RECIBIDA','** INGRESO (revisar) **'),
 ('SERV. PROF','Servicios Profesionales (honorarios)'),('HONORARIO','Servicios Profesionales (honorarios)'),
 ('COMPRA ESTABLECIMIENTO','** Compra establecimiento (revisar) **'),('PAGO DIRECTO','** Pago servicio (revisar) **'),
 ('PAGO SERVICIO','** Pago servicio (revisar) **'),('DB AUTOMATICO','** Debito automatico (revisar) **'),('TRANSFERENCIA','** Transferencia (revisar) **'),
]
def clasifica(desc):
    u=desc.upper()
    for kw,ru in REGLAS:
        if kw in u: return ru
    return 'SIN CLASIFICAR'

# --- parse verificado ---
movs=[]; recon=[]
for pdf in sorted(glob.glob(os.path.join(DEST,'Estado_Produbanco_2026-*.pdf'))):
    rd=PdfReader(pdf); txt=''
    for pg in rd.pages: txt+=(pg.extract_text() or '')+'\n'
    si=re.search(r'Saldo Inicial Contable:\s*([\d.,]+)',txt); sf=re.search(r'Saldo Final Contable:\s*([\d.,]+)',txt)
    si=num(si.group(1)); sf=num(sf.group(1))
    d=c=0.0; cnt=0
    for ln in txt.splitlines():
        m=LINE.match(ln.strip())
        if not m: continue
        _,mon,day,yr,_,desc,val,sign,_,_=m.groups()
        mi=MES.get(mon.capitalize())
        if not mi or yr!='2026': continue
        a=num(val); deb=a if sign=='-' else 0.0; cre=a if sign=='+' else 0.0
        movs.append({'mes':mi,'fecha':f'{int(day):02d}/{mi:02d}/{yr}','desc':desc.strip(),'deb':deb,'cre':cre,'rubro':clasifica(desc)})
        d+=deb; c+=cre; cnt+=1
    recon.append((os.path.basename(pdf), cnt, round(sf-si,2), round(c-d,2)))
assert all(abs(a-b)<0.05 for _,_,a,b in recon), "RECONCILIACION FALLA"
print('Movimientos:', len(movs), '(reconciliados OK en los 6 estados)')

wb=openpyxl.load_workbook(FLUJO)
thin=Side(style='thin',color='BFBFBF')
def bset(c,fmt=None,align='left',color='000000',bold=False,size=9,fill=None):
    c.font=Font(size=size,color=color,bold=bold); c.alignment=Alignment(horizontal=align,vertical='center')
    if fmt:c.number_format=fmt
    if fill:c.fill=PatternFill('solid',fgColor=fill)
    c.border=Border(left=thin,right=thin,top=thin,bottom=thin)

# --- recargar Conciliacion limpia (A-E + H) ---
wc=wb['Conciliacion']
for rr in range(5,5+400):  # limpiar
    for cc in (1,2,3,4,5,8): wc.cell(rr,cc).value=None
r=5
for mv in sorted(movs,key=lambda x:(x['mes'],int(x['fecha'][:2]))):
    bset(wc.cell(r,1,mv['fecha']),align='center'); bset(wc.cell(r,2,mv['desc']))
    bset(wc.cell(r,3,mv['deb'] or None),fmt='#,##0.00',align='right'); bset(wc.cell(r,4,mv['cre'] or None),fmt='#,##0.00',align='right')
    bset(wc.cell(r,5,'Produbanco'),align='center')
    bset(wc.cell(r,8,mv['rubro']),color=('C00000' if ('revisar' in mv['rubro'].lower() or 'SIN' in mv['rubro']) else '1F6FBF'))
    r+=1
wc.cell(4,8,'RUBRO (parser)').font=Font(bold=True)
last_conc=r-1

# --- Resumen Bancario ---
if 'Resumen Bancario' in wb.sheetnames: del wb['Resumen Bancario']
ws=wb.create_sheet('Resumen Bancario')
NAVY='1F3864'; BLUE='2E5496'; HDRG='BDD7EE'; GREEN='E2EFDA'; RED='FCE4E4'; GOLD='FFF2CC'
ws.merge_cells('A1:D1'); bset(ws.cell(1,1,'RESUMEN BANCARIO PRODUBANCO 2026 (cta. 02013011800)'),bold=True,size=12,color='FFFFFF',fill=NAVY,align='center')
bset(ws.cell(2,1,'Reconciliado: suma de movimientos = Saldo Final - Inicial en los 6 estados.'),size=9,align='left'); ws.merge_cells('A2:D2')
r=4
for i,h in enumerate(['MES','EGRESOS (deb)','INGRESOS (cre)','NETO']): bset(ws.cell(r,i+1,h),bold=True,fill=HDRG,align='center')
r+=1
bym=collections.defaultdict(lambda:[0.0,0.0])
for mv in movs: bym[mv['mes']][0]+=mv['deb']; bym[mv['mes']][1]+=mv['cre']
td=tc=0
for mi in sorted(bym):
    d,c=bym[mi]; td+=d; tc+=c
    bset(ws.cell(r,1,NOM[mi])); bset(ws.cell(r,2,round(d,2)),fmt='#,##0.00',align='right')
    bset(ws.cell(r,3,round(c,2)),fmt='#,##0.00',align='right'); bset(ws.cell(r,4,round(c-d,2)),fmt='#,##0.00',align='right',fill=(RED if c-d<0 else GREEN))
    r+=1
bset(ws.cell(r,1,'TOTAL ENE-JUN'),bold=True,fill=GOLD); bset(ws.cell(r,2,round(td,2)),bold=True,fmt='#,##0.00',align='right',fill=GOLD)
bset(ws.cell(r,3,round(tc,2)),bold=True,fmt='#,##0.00',align='right',fill=GOLD); bset(ws.cell(r,4,round(tc-td,2)),bold=True,fmt='#,##0.00',align='right',fill=GOLD)
r+=2
bset(ws.cell(r,1,'EGRESO POR RUBRO (autoclasificado)'),bold=True,color='FFFFFF',fill=BLUE); bset(ws.cell(r,2,''),fill=BLUE); r+=1
byr=collections.defaultdict(float)
for mv in movs: byr[mv['rubro']]+=mv['deb']
for ru,v in sorted(byr.items(),key=lambda x:-x[1]):
    if v<=0.005: continue
    bset(ws.cell(r,1,ru),color=('C00000' if ('revisar' in ru.lower() or 'SIN' in ru) else '000000'))
    bset(ws.cell(r,2,round(v,2)),fmt='#,##0.00',align='right'); r+=1
ws.column_dimensions['A'].width=40
for col in ('B','C','D'): ws.column_dimensions[col].width=15
ws.page_setup.orientation='portrait'; ws.page_setup.paperSize=9; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0

order=['Flujo 2026','Resumen Bancario','Conciliacion','Reglas']
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
wb.save(FLUJO)
print('Conciliacion filas:', last_conc-4, '| Resumen: total deb', round(td,2), 'cre', round(tc,2), 'neto', round(tc-td,2))
print('Hojas:', wb.sheetnames)
