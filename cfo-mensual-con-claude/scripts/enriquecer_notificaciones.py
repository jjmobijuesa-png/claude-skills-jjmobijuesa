# -*- coding: utf-8 -*-
"""Cruza notificaciones Produbanco (comercio/contraparte) con los movimientos del
estado de cuenta por monto+fecha, enriquece la descripcion y re-clasifica."""
import re, glob, os, json, datetime, collections
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

FLUJO=r"E:\vars\var 11-11 Finanza Integral\03 Flujo de Caja\Flujo de Efectivo 2026.xlsx"
DEST=r"E:\vars\var 11-11 Finanza Integral\10 Documentos de Respaldo\Estados de Cuenta\Produbanco 2026"
NOTIF=r"C:\Users\datos\AppData\Local\Temp\claude\C--Users-datos-Downloads\1e5f7d2c-d01a-49ef-b9a4-96feafbbf03a\scratchpad\notif_produbanco.json"
MES={'Jan':1,'Ene':1,'Enero':1,'Feb':2,'Febrero':2,'Mar':3,'Marzo':3,'Apr':4,'Abr':4,'Abril':4,'May':5,'Mayo':5,
     'Jun':6,'Junio':6,'Jul':7,'Julio':7,'Aug':8,'Ago':8,'Agosto':8,'Sep':9,'Set':9,'Septiembre':9,
     'Oct':10,'Octubre':10,'Nov':11,'Noviembre':11,'Dec':12,'Dic':12,'Diciembre':12}
num=lambda s: float(str(s).replace(',','').replace('$','').strip()) if s else 0.0
def parse_notif_date(s):
    m=re.match(r'(\d{1,2})/([A-Za-zé]+)/(\d{4})', s or '')
    if not m: return None
    d=int(m.group(1)); mo=MES.get(m.group(2).capitalize()); y=int(m.group(3))
    return datetime.date(y,mo,d) if mo else None

# --- reglas (mismas del sistema) ---
REGLAS=[
 ('TARIFA','Comisiones bancarias'),('REDONDEO','Aporte / Ahorro'),('AHORRO INCREMENTAL','Aporte / Ahorro'),
 ('FLEXIAHORRO','Aporte / Ahorro'),('AHORRAR POR AHORRAR','Aporte / Ahorro'),
 ('RETIROS ATM','** Retiro efectivo (revisar) **'),('RETIRO DE EFECTIVO','** Retiro efectivo (revisar) **'),
 ('SANA SANA','Medicina'),('FYBECA','Medicina'),('FARMACIA','Medicina'),('PHARMACY','Medicina'),('CRUZ AZUL','Medicina'),
 ('MEDICITY','Medicina'),('DIFARE','Medicina'),('CLINICA','Medicina'),('HOSPITAL','Medicina'),('LABORATORIO','Medicina'),('OPTICA','Medicina'),
 ('CNEL','Servicios Basicos'),('AGUA POTABLE','Servicios Basicos'),('EMAPA','Servicios Basicos'),
 ('CNT','Internet CNT'),('CLARO','Movil Claro'),('MOVISTAR','Movil Claro'),('NETLIFE','Internet CNT'),
 ('PRIMAX','Combustible'),('PETROECUADOR','Combustible'),('MASGAS','Combustible'),('PRODUGAS','Combustible'),('GASOLINERA','Combustible'),('TERPEL','Combustible'),
 ('ROSADO','Alimentos vivienda'),('SUPERMAXI','Alimentos vivienda'),('MEGAMAXI','Alimentos vivienda'),('COMISARIATO','Alimentos vivienda'),
 ('SANTA MARIA','Alimentos vivienda'),('GRAN AKI','Alimentos vivienda'),('TIA','Alimentos vivienda'),('CARNE','Alimentos vivienda'),('PANADERIA','Alimentos vivienda'),('TERRAMONTE','Alimentos vivienda'),
 ('SUPERDEPORTE','Ropa'),('DE PRATI','Ropa'),('ETAFASHION','Ropa'),('MARATHON','Ropa'),
 ('TECNICENTRO','Vehiculo (seguro-aceite)'),('LUBRICANT','Vehiculo (seguro-aceite)'),('LLANTA','Vehiculo (seguro-aceite)'),
 ('PARQUEO','Parqueo'),('PARKING','Parqueo'),('PEAJE','Peajes/Viajes'),('TAXI','Peajes/Viajes'),
 ('IESS','IESS'),('DIVIDENDO','Servicio de Deuda'),('PRESTAMO','Servicio de Deuda'),('AMORTIZA','Servicio de Deuda'),
 ('MUTUALISTA','Servicio de Deuda'),('LA NUESTRA','Servicio de Deuda'),('MOBIJUESA','Servicio de Deuda'),('COOPERATIVA','Servicio de Deuda'),
 ('NIPPONFLEX','Oficina - Nipponflex'),('DONACION','Donacion'),('DIEZMO','Donacion'),('IGLESIA','Donacion'),
 ('DEPOSITO','** INGRESO (revisar) **'),('TRANSFERENCIA RECIBIDA','** INGRESO (revisar) **'),
 ('SERV. PROF','Servicios Profesionales (honorarios)'),('HONORARIO','Servicios Profesionales (honorarios)'),
 ('INTERBANCARIA','** Transferencia (revisar) **'),('TRANSFERENCIA CTAS TERCEROS','** Transferencia (revisar) **'),
 ('TRANSFERENCIA ENTRE CTAS','** Transferencia (revisar) **'),('TRANSFERENCIA CTAS PROPIAS','** Transferencia (revisar) **'),
 ('COMPRA ESTABLECIMIENTO','** Compra establecimiento (revisar) **'),('PAGO DIRECTO','** Pago servicio (revisar) **'),
 ('PAGO SERVICIO','** Pago servicio (revisar) **'),('DB AUTOMATICO','** Debito automatico (revisar) **'),('TRANSFERENCIA','** Transferencia (revisar) **'),
]
def clasifica(desc):
    u=(desc or '').upper()
    for kw,ru in REGLAS:
        if kw in u: return ru
    return 'SIN CLASIFICAR'

# --- cargar notificaciones (con tipo, para cruce preciso) ---
def notif_type(nt):
    if nt.get('establecimiento'): return 'consumo'
    if nt.get('contacto') or nt.get('banco_destino'): return 'transf_out'
    if nt.get('deposito') or re.search(r'(Recibida|Ingresada|Acreditada|Dep)', nt.get('subj',''), re.I): return 'ingreso'
    return 'otro'
notifs=json.loads(open(NOTIF,encoding='utf-8').read())
idx=collections.defaultdict(list)   # monto -> [(fecha,label,tipo)]
for nt in notifs:
    val=num(nt.get('valor') or nt.get('deposito'))
    if val<=0: continue
    d=parse_notif_date(nt.get('fecha'))
    label = nt.get('establecimiento') or nt.get('contacto') or ''
    if nt.get('banco_destino'): label += (' ['+nt['banco_destino']+']')
    if nt.get('descripcion') and nt.get('descripcion').lower() not in ('','dnd'): label += (' - '+nt['descripcion'])
    if not label and nt.get('deposito'): label='Deposito '+(nt.get('canal') or '')
    idx[round(val,2)].append((d,label.strip(),notif_type(nt)))
print('Notificaciones con monto:', sum(len(v) for v in idx.values()))

def mov_type(desc, sign):
    u=desc.upper()
    if 'COMPRA ESTABLEC' in u: return 'consumo'
    if 'RETIRO' in u: return 'retiro'
    if 'TRANSFER' in u and sign=='-': return 'transf_out'
    if sign=='+': return 'ingreso'
    return 'otro'

# --- re-parsear movimientos y enriquecer ---
LINE=re.compile(r'^([A-ZÁÉÍÓÚÑ]{3,})\s+([A-Za-z]{3})\s+(\d{1,2})\s+(\d{4})\s+(\S+)\s+(.+?)\s+([\d.,]+)([+-])\s+([\d.,]+)\s+([\d.,]+)\s*$')
MESB={'Jan':1,'Ene':1,'Feb':2,'Mar':3,'Apr':4,'Abr':4,'May':5,'Jun':6,'Jul':7,'Aug':8,'Ago':8,'Sep':9,'Set':9,'Oct':10,'Nov':11,'Dec':12,'Dic':12}
movs=[]; enriched=0
for pdf in sorted(glob.glob(os.path.join(DEST,'Estado_Produbanco_2026-*.pdf'))):
    rd=PdfReader(pdf); txt=''
    for pg in rd.pages: txt+=(pg.extract_text() or '')+'\n'
    for ln in txt.splitlines():
        m=LINE.match(ln.strip())
        if not m: continue
        _,mon,day,yr,_,desc,val,sign,_,_=m.groups()
        mi=MESB.get(mon.capitalize())
        if not mi or yr!='2026': continue
        amount=num(val); fecha=datetime.date(2026,mi,int(day))
        mt=mov_type(desc, sign)
        extra=''; best=None
        for d,label,nt_ in idx.get(round(amount,2),[]):
            if not label or not d: continue
            if abs((d-fecha).days)>2: continue
            # cruce por tipo compatible
            if mt=='consumo' and nt_ not in ('consumo','transf_out'): continue
            if mt=='transf_out' and nt_!='transf_out': continue
            if mt=='ingreso' and nt_!='ingreso': continue
            if mt in ('retiro','otro'): continue
            best=label; break
        if best:
            extra=' | '+best; enriched+=1
        full=desc.strip()+extra
        movs.append({'mes':mi,'fecha':fecha,'desc':full,'deb':amount if sign=='-' else 0.0,
                     'cre':amount if sign=='+' else 0.0,'rubro':clasifica(full)})
print('Movimientos:',len(movs),'| enriquecidos con comercio:',enriched)

# --- recargar Conciliacion ---
wb=openpyxl.load_workbook(FLUJO); wc=wb['Conciliacion']
thin=Side(style='thin',color='BFBFBF')
def bset(c,fmt=None,align='left',color='000000'):
    c.font=Font(size=9,color=color); c.alignment=Alignment(horizontal=align,vertical='center')
    if fmt:c.number_format=fmt
    c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
for rr in range(5,5+400):
    for cc in (1,2,3,4,5,8): wc.cell(rr,cc).value=None
r=5
for mv in sorted(movs,key=lambda x:(x['fecha'])):
    bset(wc.cell(r,1,mv['fecha']),fmt='dd/mm/yyyy',align='center'); bset(wc.cell(r,2,mv['desc']))
    bset(wc.cell(r,3,mv['deb'] or None),fmt='#,##0.00',align='right'); bset(wc.cell(r,4,mv['cre'] or None),fmt='#,##0.00',align='right')
    bset(wc.cell(r,5,'Produbanco'),align='center')
    bset(wc.cell(r,8,mv['rubro']),color=('C00000' if ('revisar' in mv['rubro'].lower() or 'SIN' in mv['rubro']) else '1F6FBF'))
    r+=1

# --- Resumen Bancario actualizado ---
byr=collections.defaultdict(float)
for mv in movs: byr[mv['rubro']]+=mv['deb']
ws=wb['Resumen Bancario']
# reescribir bloque rubros (buscar fila 'EGRESO POR RUBRO')
start=None
for rr in range(1,60):
    if ws.cell(rr,1).value and 'EGRESO POR RUBRO' in str(ws.cell(rr,1).value): start=rr+1; break
if start:
    for rr in range(start, start+40):
        ws.cell(rr,1).value=None; ws.cell(rr,2).value=None
    rr=start
    for ru,v in sorted(byr.items(),key=lambda x:-x[1]):
        if v<=0.005: continue
        c=ws.cell(rr,1,ru); c.font=Font(size=9,color=('C00000' if ('revisar' in ru.lower() or 'SIN' in ru) else '000000'))
        c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        c2=ws.cell(rr,2,round(v,2)); c2.number_format='#,##0.00'; c2.alignment=Alignment(horizontal='right'); c2.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        rr+=1
wb.save(FLUJO)
# resumen enriquecimiento
rev=sum(v for k,v in byr.items() if 'revisar' in k.lower() or 'SIN' in k)
tot=sum(byr.values())
print('Egreso total %.2f | aun en revisar %.2f (%.0f%%)'%(tot,rev,100*rev/tot if tot else 0))
print('EGRESO POR RUBRO tras enriquecer:')
for ru,v in sorted(byr.items(),key=lambda x:-x[1]):
    if v>0.005: print('  %9.2f  %s'%(v,ru))
