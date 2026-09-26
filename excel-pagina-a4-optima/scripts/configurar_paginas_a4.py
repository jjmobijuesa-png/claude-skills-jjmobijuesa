# -*- coding: utf-8 -*-
"""
Configura la CONFIGURACIÓN DE PÁGINA de cada hoja de un .xlsx para que el
contenido ocupe la MAYOR superficie posible de una hoja A4, eligiendo
automáticamente orientación (horizontal / vertical) y el % de escala.

No solo "que quepa": calcula la escala para que el contenido LLENE la A4.

Uso:
    python configurar_paginas_a4.py "archivo.xlsx" [salida.xlsx]
    (si no se da salida, sobreescribe el mismo archivo — hace backup .bak)

Reglas:
  - papel A4 (paperSize=9).
  - orientación = la que permita MAYOR factor de llenado.
  - escala = round(min(anchoUtil/anchoContenido, altoUtil/altoContenido)*100),
    acotada a [40, 250]% (evita miniaturizar o inflar en exceso).
  - márgenes estrechos (0.5 cm) + centrado horizontal y vertical.
  - fitToPage se desactiva para respetar la escala explícita calculada.
  - si el contenido es MÁS ancho/alto que una A4 aun al mínimo, se usa
    fitToWidth=1 (todas las columnas en 1 página de ancho) como respaldo.
"""
import sys, shutil
from pathlib import Path
import openpyxl
from openpyxl.worksheet.page import PageMargins

# A4 en puntos (1 pulg = 72 pt). A4 = 210 x 297 mm.
MM = 72 / 25.4
A4_W_MM, A4_H_MM = 210.0, 297.0
MARGIN_CM = 0.5           # márgenes estrechos
MARGIN_PT = MARGIN_CM * 10 * MM
SCALE_MIN, SCALE_MAX = 40, 250

def col_width_pt(w):
    # openpyxl width ≈ nº de caracteres; aprox 7 px/char a 96dpi -> pt = char*7*72/96
    if w is None: w = 8.43         # ancho por defecto de Excel
    px = w * 7 + 5
    return px * 72 / 96

def row_height_pt(h):
    return h if h else 15.0        # alto por defecto ≈ 15 pt

def content_size(ws):
    # rango usado
    if ws.max_row < 1 or ws.max_column < 1:
        return 0.0, 0.0
    total_w = 0.0
    for c in range(1, ws.max_column + 1):
        letter = openpyxl.utils.get_column_letter(c)
        dim = ws.column_dimensions.get(letter)
        total_w += col_width_pt(dim.width if dim and dim.width else None)
    total_h = 0.0
    for r in range(1, ws.max_row + 1):
        dim = ws.row_dimensions.get(r)
        total_h += row_height_pt(dim.height if dim and dim.height else None)
    return total_w, total_h

def fill_scale(content_w, content_h, page_w_mm, page_h_mm):
    usable_w = page_w_mm * MM - 2 * MARGIN_PT
    usable_h = page_h_mm * MM - 2 * MARGIN_PT
    if content_w <= 0 or content_h <= 0:
        return 100
    s = min(usable_w / content_w, usable_h / content_h) * 100
    return s

def configure(path_in, path_out=None):
    path_in = Path(path_in)
    if path_out is None:
        path_out = path_in
        bak = path_in.with_suffix(path_in.suffix + ".bak")
        if not bak.exists():
            shutil.copy(path_in, bak)
            print(f"  backup: {bak.name}")
    wb = openpyxl.load_workbook(path_in)
    for ws in wb.worksheets:
        cw, ch = content_size(ws)
        # probar ambas orientaciones, elegir la de mayor llenado
        s_port = fill_scale(cw, ch, A4_W_MM, A4_H_MM)   # vertical
        s_land = fill_scale(cw, ch, A4_H_MM, A4_W_MM)   # horizontal
        if s_land >= s_port:
            orient = "landscape"; scale = s_land
        else:
            orient = "portrait";  scale = s_port

        ws.page_setup.orientation = orient
        ws.page_setup.paperSize = 9  # A4
        ws.page_setup.fitToWidth = 0
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=False)
        ws.page_margins = PageMargins(left=MARGIN_CM/2.54, right=MARGIN_CM/2.54,
                                      top=MARGIN_CM/2.54, bottom=MARGIN_CM/2.54,
                                      header=0.2, footer=0.2)
        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered = True

        clamped = max(SCALE_MIN, min(SCALE_MAX, round(scale)))
        if scale < SCALE_MIN:
            # el contenido no cabe ni al mínimo: mejor "ajustar a 1 pág de ancho"
            ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            modo = "fit-ancho (contenido grande)"
        else:
            ws.page_setup.scale = clamped
            modo = f"escala {clamped}%"
        print(f"  [{ws.title}] {orient} · {modo}  (contenido {cw:.0f}x{ch:.0f} pt)")
    wb.save(path_out)
    print(f"OK -> {path_out}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    configure(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
