# -*- coding: utf-8 -*-
"""Conversor Markdown -> Word (.docx) con soporte de tablas, encabezados,
viñetas, casillas, negritas/cursivas, citas y reglas horizontales."""
import re, sys, os, io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY  = RGBColor(0x0F, 0x2E, 0x54)
STEEL = RGBColor(0x2C, 0x5F, 0x8A)
GREY  = RGBColor(0x55, 0x55, 0x55)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    run.font.size = Pt(9); run.font.color.rgb = GREY

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def add_inline(par, text):
    """Escribe texto con formato inline (negrita, cursiva, código, enlaces)."""
    text = text.replace('\\|', '|')
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        elif tok.startswith('[') and '](' in tok:
            label = tok[1:tok.index('](')]
            r = par.add_run(label); r.font.color.rgb = STEEL; r.underline = True
        else:
            par.add_run(tok)

def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    # respeta escapes \|
    parts = re.split(r'(?<!\\)\|', line)
    return [p.strip() for p in parts]

def is_sep(line):
    return bool(re.match(r'^\s*\|?[\s:\-\|]+\|[\s:\-\|]*$', line)) and '-' in line

def convert(md_path, out_path):
    with io.open(md_path, encoding='utf-8') as fh:
        lines = fh.read().split('\n')

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.12

    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    add_page_numbers(sec)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()

        # --- tabla ---
        if s.startswith('|') and i + 1 < n and is_sep(lines[i+1]):
            headers = split_row(s)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            ncols = len(headers)
            t = doc.add_table(rows=1, cols=ncols)
            t.style = 'Light Grid Accent 1'
            for c, htxt in enumerate(headers):
                cell = t.rows[0].cells[c]
                cell.text = ''
                p = cell.paragraphs[0]
                r = p.add_run(re.sub(r'\*\*', '', htxt))
                r.bold = True; r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_bg(cell, '0F2E54')
            for row in rows:
                cells = t.add_row().cells
                for c in range(ncols):
                    val = row[c] if c < len(row) else ''
                    cells[c].text = ''
                    p = cells[c].paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    add_inline(p, val)
                    for rr in p.runs: rr.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # --- encabezados ---
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            level = len(m.group(1)); txt = m.group(2).strip()
            if level == 1:
                p = doc.add_heading(level=1)
                r = p.add_run(txt); r.font.size = Pt(16); r.font.color.rgb = NAVY; r.font.name='Calibri'
            elif level == 2:
                p = doc.add_heading(level=2)
                r = p.add_run(txt); r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name='Calibri'
            elif level == 3:
                p = doc.add_heading(level=3)
                r = p.add_run(txt); r.font.size = Pt(11.5); r.font.color.rgb = STEEL; r.font.name='Calibri'
            else:
                p = doc.add_heading(level=4)
                r = p.add_run(txt); r.font.size = Pt(10.5); r.font.color.rgb = STEEL; r.font.name='Calibri'
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # --- regla horizontal ---
        if re.match(r'^\s*(-{3,}|_{3,}|\*{3,})\s*$', s):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'BBBBBB')
            pbdr.append(bottom); pPr.append(pbdr)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # --- cita ---
        if s.startswith('>'):
            block = []
            while i < n and lines[i].strip().startswith('>'):
                block.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, ' '.join(x for x in block if x))
            for r in p.runs:
                r.italic = True; r.font.color.rgb = STEEL; r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(7)
            continue

        # --- bloque de código ---
        if s.startswith('```'):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            for cl in code:
                p = doc.add_paragraph()
                r = p.add_run(cl)
                r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x30,0x30,0x30)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # --- lista numerada ---
        m = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_inline(p, m.group(2))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # --- viñeta / casilla ---
        m = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if m:
            indent = len(m.group(1)); body = m.group(2)
            chk = re.match(r'^\[( |x|X)\]\s*(.*)$', body)
            if chk:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75 + 0.5*(indent//2))
                mark = p.add_run(('☒ ' if chk.group(1).lower()=='x' else '☐ '))
                mark.font.size = Pt(11)
                add_inline(p, chk.group(2))
            else:
                style = 'List Bullet' if indent < 2 else 'List Bullet 2'
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph(style='List Bullet')
                add_inline(p, body)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # --- vacío ---
        if not s:
            i += 1
            continue

        # --- párrafo normal ---
        p = doc.add_paragraph()
        add_inline(p, s)
        p.paragraph_format.space_after = Pt(6)
        i += 1

    doc.save(out_path)
    return out_path

if __name__ == '__main__':
    for md in sys.argv[1:]:
        out = os.path.splitext(md)[0] + '.docx'
        convert(md, out)
        print('OK ->', os.path.basename(out))
