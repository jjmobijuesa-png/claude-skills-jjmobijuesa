# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0F, 0x2E, 0x54)
STEEL = RGBColor(0x2C, 0x5F, 0x8A)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---- Base styles ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_page_numbers(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    run.font.size = Pt(9); run.font.color.rgb = GREY

def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text); r.font.color.rgb = NAVY; r.font.size = Pt(15); r.font.name='Calibri'
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    return p

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text); r.font.color.rgb = STEEL; r.font.size = Pt(12.5); r.font.name='Calibri'
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
    return p

def para(text, italic=False, bold=False, size=11, color=None, align=None, after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic=italic; r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_lead:
        r = p.add_run(bold_lead); r.bold=True
        p.add_run(text)
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True; run.font.size = Pt(10); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        set_cell_bg(hdr[i], '0F2E54')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

# ================= PORTADA =================
sec = doc.sections[0]
sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

for _ in range(3): doc.add_paragraph()
para('FUNDACIÓN BLOCKCHAIN SOBERANA DEL ECUADOR', bold=True, size=13, color=STEEL, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para('EcuaLedger Soberana', bold=True, size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

para('PROYECTO DE CALIFICACIÓN ANTE ENTIDAD FINANCIERA MULTILATERAL', bold=True, size=20, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para('Postulación a BID Lab y al Fondo de Impacto VELA de CAF', size=14, color=STEEL, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para('Infraestructura Pública Digital Blockchain Soberana para la Inclusión Económica Digital y la Tokenización de Activos del Mundo Real (RWA) en el Ecuador', italic=True, size=11.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)

para('PRIMER BORRADOR — Versión 1.0', bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para('Quevedo, Provincia de Los Ríos — Ecuador  ·  Julio de 2026', size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para('Clasificación: Uso Restringido — Documento de trabajo', italic=True, size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

add_page_numbers(sec)
doc.add_page_break()

# ================= 0. NOTA PRELIMINAR =================
h1('Nota preliminar')
para('El presente documento constituye el primer borrador del proyecto con el que la Fundación Blockchain Soberana del Ecuador (en adelante, «la Fundación» o «FBSE») busca la calificación, validación y asignación de fondos por parte de las entidades financieras multilaterales que auspician la inversión de impacto y el capital emprendedor en América Latina y el Caribe: el Banco Interamericano de Desarrollo, a través de su laboratorio de innovación BID Lab, y el Banco de Desarrollo de América Latina y el Caribe (CAF), a través de su Fondo de Impacto VELA.')
para('A diferencia del antecedente del proyecto —centrado en CAF y en la República del Paraguay como captador regional del modelo Legaledger, ya regulado por las Leyes 6822/21 y 7572/25—, este borrador desplaza el foco hacia los organismos que califican el proyecto, lo validan y le asignan los fondos de manera directa, y adapta la propuesta ecuatoriana (EcuaLedger Soberana) al lenguaje y a los criterios de elegibilidad de cada fondo.')
para('Se trata de un documento vivo: las cifras, citas normativas y estados institucionales señalados deben verificarse en su redacción vigente antes de cualquier presentación formal, conforme a las advertencias de precisión de la Sección 8.', italic=True, color=GREY)

# ================= 1. RESUMEN EJECUTIVO =================
h1('1. Resumen ejecutivo')
para('EcuaLedger Soberana es la primera red de registros distribuidos (blockchain) de tercera generación con soberanía nacional propuesta para la República del Ecuador. Inspirada en el precedente paraguayo Legaledger —caso de uso en evaluación ante el Comité Técnico ISO/TC 307—, no es un activo especulativo, sino la infraestructura pública digital que blinda la seguridad jurídica, optimiza la eficiencia estatal y habilita la tokenización de activos del mundo real (RWA) en una economía dolarizada.')
para('La Fundación busca capital de impacto para desplegar el Laboratorio Alfa y el Sandbox Regulatorio que anteceden a la TestNet y a la MainNet soberanas, con foco en inclusión financiera digital y trazabilidad productiva rural. Para ello se dirige a dos ventanillas complementarias:')
bullet(' invierte en fondos de capital emprendedor (venture capital) de la región; su puerta natural para una infraestructura blockchain es el ecosistema LACChain del propio BID y, en paralelo, sus instrumentos de cooperación técnica.', bold_lead='BID Lab —')
bullet(' primer fondo de inversión de impacto regional de CAF (US$100–150 millones, con Sonen Capital y Fondo de Fondos), que invierte en empresas, fondos y soluciones de alto impacto con retorno medible; el encaje de EcuaLedger es la inclusión financiera y la agricultura regenerativa.', bold_lead='CAF — Fondo VELA —')
para('Hallazgo central de este borrador: ninguno de los dos fondos es una donación asistencial. Ambos invierten para obtener impacto con retorno. La convocatoria de venture capital del BID Lab, además, financia gestoras de fondos y no fundaciones ni proyectos de forma directa. En consecuencia, la Fundación debe presentar el proyecto como una oportunidad de inversión escalable y definir el vehículo invertible correspondiente, según se detalla en la Sección 4.', bold=True)

# ================= 2. EL PROYECTO =================
h1('2. El proyecto: EcuaLedger Soberana')
h2('2.1 Qué es')
para('EcuaLedger es una red blockchain soberana, pública-permisionada y sin comisión por transacción (gas fee), con gobernanza multisectorial en la que el Estado, el sector privado, la academia y las cooperativas controlan cada uno el 25 % de los nodos validadores. Su propósito exclusivo es dar certeza jurídica a las transacciones y actos jurídicos, públicos o privados, tramitados electrónicamente. Su naturaleza es equiparable a un registro notarial en formato digital distribuido y supervisado.')
para('Arquitectura de referencia: doble motor Hyperledger Fabric + Hyperledger Besu, módulo de seguridad por hardware (HSM) con certificación NIST, consenso por Prueba de Autoridad (PoA), firmas con sellado de tiempo y contratos inteligentes con auditoría obligatoria; preparada para cifrado post-cuántico.')

h2('2.2 Los dos pilares normativos')
para('Para que los tokens emitidos representen un derecho jurídicamente exigible, la Fundación promueve un marco homólogo al paraguayo:')
bullet(' consagra la equivalencia funcional y el no repudio: un documento, firma, sello o token en EcuaLedger tiene igual o superior valor jurídico y probatorio que su soporte físico. Homóloga de la Ley N.º 6822/21 del Paraguay.', bold_lead='Pilar 1 — Ley Orgánica de la Fe Pública Digital (LOFPD): ')
bullet(' regula la emisión de activos del mundo real (RWA), el crowdfunding y el registro de valores mediante tecnología de registros distribuidos (DLT). Homóloga de la Ley N.º 7572/25 del Paraguay.', bold_lead='Pilar 2 — Reforma del Mercado de Valores y Comercialización Tokenizada (RWA): ')
para('Anclaje en Ecuador: la seguridad jurídica que Paraguay obtuvo por ley, Ecuador la tiene a nivel constitucional (art. 82), lo que permite iniciar el Sandbox y el Alfa Lab sobre el marco vigente, sin reforma legislativa previa, y generar la evidencia que precede a la ley modelo.')

h2('2.3 Innovación jurídica central')
para('EcuaLedger separa, a nivel de protocolo, los dos atributos del dominio del derecho civil continental: la Posesión (ius possidendi) y la Propiedad (ius disponendi). Ello permite separar el financiamiento del uso: un productor puede tokenizar su ganado o su cosecha transfiriendo la posesión al acreedor como garantía real, manteniendo la propiedad, sin escritura pública ni intermediario notarial —habilitando microcrédito con garantía verificable para pequeños productores.')

h2('2.4 Hoja de ruta de despliegue')
para('Laboratorio Alfa  →  Sandbox Regulatorio  →  Catálogo de Soluciones Certificadas  →  TestNet y MainNet soberanas.', bold=True)

# ================= 3. LOS FONDOS =================
h1('3. Los fondos auspiciantes: especificaciones')
h2('3.1 BID Lab — Convocatoria de Fondos de Venture Capital')
table(
    ['Campo', 'Especificación'],
    [
     ['Organismo', 'BID Lab (Laboratorio de Innovación del Grupo BID)'],
     ['Objetivo', 'Invertir en fondos de VC que cierren brechas de financiamiento de empresas en etapa temprana con impacto'],
     ['Naturaleza', 'Invierte EN LOS FONDOS (gestoras), NO en empresas, fundaciones ni proyectos de forma directa'],
     ['Quién postula', 'Firmas gestoras de fondos; aceleradoras y venture builders que evolucionan a vehículos de inversión'],
     ['Sectores', 'Fintech, climate tech, agtech, healthtech, edtech; empleo, productividad e inclusión'],
     ['Geografía', 'América Latina y el Caribe (Ecuador no excluido)'],
     ['Fecha límite', '30 de septiembre de 2026 (inversiones esperadas en 2027)'],
     ['URL oficial', 'iadb.org/en/home/calls-proposals/venture-capital-funds'],
     ['Vía natural para EcuaLedger', 'LACChain (ecosistema blockchain del BID) + instrumentos de cooperación técnica'],
    ],
    widths=[4.0, 12.0]
)
h2('3.2 CAF — Fondo de Impacto VELA')
table(
    ['Campo', 'Especificación'],
    [
     ['Organismo', 'CAF — Banco de Desarrollo de América Latina y el Caribe'],
     ['Instrumento', 'Fondo de Impacto VELA — primer fondo de inversión de impacto regional de CAF'],
     ['Gestores', 'CAF (inversor ancla) con Sonen Capital LLC y Fondo de Fondos'],
     ['Tamaño objetivo', 'US$100–150 millones hacia empresas, fondos y soluciones de alto potencial'],
     ['Sectores', 'Biodiversidad, agricultura regenerativa, inclusión financiera, salud, educación'],
     ['A quién financia', 'Empresas, fondos y soluciones con impacto medible y retorno (no es donación)'],
     ['Lanzamiento', 'Virtual, 21 de julio de 2026'],
     ['URL oficial', 'caf.com/es/actualidad/eventos/participe-en-el-lanzamiento-virtual-del-fondo-de-impacto-vela/'],
     ['Encaje EcuaLedger', 'Inclusión financiera digital rural + trazabilidad de agricultura regenerativa'],
    ],
    widths=[4.0, 12.0]
)

# ================= 4. ANALISIS DE ENCAJE Y CUMPLIMIENTO =================
h1('4. Análisis de encaje y cumplimiento')
para('Ambos fondos invierten para obtener impacto con retorno. Todo el material debe leerse como oportunidad de inversión escalable —problema, solución, población objetivo, teoría de cambio, métricas de impacto, monto solicitado y, de manera ineludible, retorno o mecanismo de sostenibilidad—, no como solicitud asistencial.')

h2('4.1 Rutas de encaje ante BID Lab')
para('La convocatoria de venture capital financia gestoras de fondos; por tanto, la Fundación no es postulante directo de esa ventanilla. Sus rutas, en orden de viabilidad, son:')
table(
    ['Ruta', 'Descripción', 'Viabilidad'],
    [
     ['B-1  LACChain', 'Posicionar EcuaLedger como nodo/infraestructura nacional interoperable en el ecosistema blockchain del BID Lab; asistencia técnica y visibilidad sin constituir fondo', 'Alta (primer paso)'],
     ['B-2  Pipeline de un fondo', 'EcuaLedger como caso de pipeline (fintech/agtech) de una gestora que sí postule a la convocatoria de VC', 'Media'],
     ['B-3  Vehículo propio', 'FBSE ancla o co-constituye un fondo de VC de impacto digital elegible ante BID Lab', 'Baja-media'],
     ['B-4  Otra ventanilla BID', 'Instrumentos de donación / cooperación técnica / equity a proyectos de innovación para financiar el Alfa Lab', 'Media'],
    ],
    widths=[3.0, 10.5, 2.5]
)
para('Recomendación: abrir LACChain (B-1) de inmediato y explorar B-4 para financiamiento directo del Alfa Lab; reservar B-2 y B-3 como estrategia de mediano plazo.', bold=True)

h2('4.2 Matriz de elegibilidad ante el Fondo VELA')
table(
    ['Criterio VELA', 'Estado', 'Evidencia / brecha'],
    [
     ['Sector de impacto declarado', 'Cumple (alto)', 'Inclusión financiera y agricultura regenerativa: identidad digital rural, trazabilidad agro, microcrédito con garantía real'],
     ['Impacto medible y rendición de cuentas', 'Cumple', 'Matriz de indicadores existente; la blockchain genera datos verificables en tiempo real'],
     ['Retorno / sostenibilidad', 'A fortalecer', 'Modelo de ingresos definido (formación, licenciamiento SaaS, microtarifas, asistencia técnica); convertir en proyección de retorno'],
     ['Escalamiento', 'Cumple', 'Modelo replicable en Colombia, Perú, Bolivia y Centroamérica'],
     ['Gobernanza', 'Cumple', 'Multisectorial 25/25/25/25; auditoría externa de smart contracts; separación de roles'],
     ['Vehículo invertible', 'Brecha', 'FBSE es fundación sin fines de lucro; definir solución licenciable, spin-off o co-inversión con retorno'],
     ['AML/CFT, KYC, datos', 'Cumple', 'Cumplimiento por diseño; protección de datos (LOPDP 2021)'],
    ],
    widths=[3.5, 2.2, 10.3]
)
para('Brecha principal (VELA): definir cómo se invierte en EcuaLedger —solución licenciable con ingresos, spin-off empresarial o co-inversión en el vehículo que despliega la infraestructura—. Sin vehículo invertible, VELA no puede asignar capital.', bold=True)

# ================= 5. IMPACTO =================
h1('5. Teoría de cambio e impacto')
para('El proyecto corrige tres fallas de mercado verificables: (i) falla de información —los pequeños productores carecen de registros verificables que los excluyen del sistema financiero—; (ii) falla de acceso a infraestructura digital de calidad en el territorio rural; y (iii) falla de capacidades técnicas en cooperativas y gobiernos locales.')
para('Objetivos específicos e indicadores (metas a Fase 3):')
table(
    ['Indicador', 'Meta'],
    [
     ['Nodos blockchain soberanos operativos', '6 nodos'],
     ['Cooperativas conectadas a identidad digital y trazabilidad', '40 cooperativas'],
     ['Personas capacitadas y certificadas', '2.500 personas'],
     ['Transacciones productivas registradas en cadena', '150.000 registros'],
     ['Reducción de costos de intermediación', '18–25 %'],
     ['Modelo replicable documentado + convenios regionales', '1 modelo + 3 convenios'],
    ],
    widths=[10.5, 5.5]
)
para('ODS prioritarios: 1 (fin de la pobreza), 2 (hambre cero), 4 (educación), 8 (trabajo decente), 9 (industria e innovación), 10 (reducción de desigualdades) y 17 (alianzas).')

# ================= 6. ARQUITECTURA FINANCIERA =================
h1('6. Arquitectura financiera y tesis de inversión')
para('El proyecto opera bajo una lógica de inversión de impacto escalonada, con capital multilateral como ancla que desbloquea el cofinanciamiento local y la co-inversión privada. La estructura híbrida recomendada combina un tramo de capital semilla / asistencia técnica, un tramo de inversión de impacto y cofinanciamiento local en especie.')
table(
    ['Fuente', 'Rol', 'Monto (USD)', '%'],
    [
     ['Capital multilateral de impacto (CAF VELA / BID)', 'Inversión ancla, Fases 1–2', '1.800.000', '45 %'],
     ['SENESCYT / SENPLADES', 'Cofinanciamiento ciencia e innovación', '600.000', '15 %'],
     ['GAD provinciales y municipales', 'Infraestructura local y operación', '400.000', '10 %'],
     ['CNT EP', 'Conectividad (aporte en especie)', '300.000', '7,5 %'],
     ['UTEQ y universidades socias', 'I+D, formación y espacio físico', '300.000', '7,5 %'],
     ['Fondos privados / impacto', 'Capital complementario Fases 3–4', '600.000', '15 %'],
     ['TOTAL PROYECTO', '', '4.000.000', '100 %'],
    ],
    widths=[6.5, 5.0, 2.7, 1.8]
)
para('Sostenibilidad post-inversión (base del retorno para el inversor de impacto): servicios de formación y certificación; licenciamiento de la plataforma de trazabilidad e identidad bajo modelo SaaS; asistencia técnica reembolsable; y microtarifas por transacción de trazabilidad registrada en cadena.')
para('Nota: el monto de US$1,8 M es una estimación de referencia; debe validarse con el área correspondiente del fondo antes de la presentación formal, y prepararse una versión reducida (Fases 1–2) para facilitar aprobaciones rápidas.', italic=True, color=GREY)

# ================= 7. REQUISITOS INSTITUCIONALES + SECUENCIA =================
h1('7. Requisitos institucionales y secuencia de acción')
h2('7.1 Requisitos institucionales de la Fundación')
table(
    ['Documento / condición', 'Acción'],
    [
     ['Personería jurídica y estatuto vigente con fines compatibles', 'Confirmar / reformar objeto (blockchain, innovación, inclusión)'],
     ['Representante legal designado y vigente', 'Confirmar registro de directiva'],
     ['Cumplimiento UAFE (prevención de lavado)', 'Verificar registro'],
     ['Cuenta bancaria institucional operativa', 'Confirmar'],
     ['Convenio / carta de intención con UTEQ', 'Gestionar (plantilla de MdE disponible)'],
     ['Acercamiento CNT / GAD / SENESCYT', 'Gestionar cartas de intención'],
     ['Capacidad de contrapartida en especie (20–30 %)', 'Documentar'],
    ],
    widths=[9.5, 6.5]
)
h2('7.2 Secuencia recomendada (90 días)')
bullet('Semanas 1–2: abrir canal LACChain (BID) con ficha técnica de EcuaLedger; registrar interés en VELA ante CAF / Sonen Capital / Fondo de Fondos.')
bullet('Semanas 2–4: redactar la Nota de Inversión de Impacto (2–4 páginas) reutilizando el paquete existente y añadiendo la sección de retorno / vehículo invertible.')
bullet('Semanas 3–6: cerrar cartas de intención (UTEQ, CNT, GAD, SENESCYT) y confirmar la documentación institucional de la Fundación.')
bullet('Semanas 4–8: definir el vehículo invertible para VELA y evaluar la ventanilla de cooperación técnica del BID Lab para el Alfa Lab.')
bullet('Antes del 30-sep-2026: si se optó por las rutas B-2/B-3, co-presentar a la convocatoria de venture capital del BID Lab con la gestora aliada.')

# ================= 8. ADVERTENCIAS =================
h1('8. Advertencias de precisión')
bullet('El estado ante ISO/TC 307 es una evaluación en curso, no una aprobación: redactar siempre en condicional.')
bullet('Las citas normativas (Paraguay 6822/21 y 7572/25; Ecuador: Constitución, LOPDP, Ley Fintech, COMYF/SEPS) deben verificarse en su redacción vigente antes del envío.')
bullet('Diferenciar el Fondo VELA (US$100–150 M, impacto) de un eventual fondo de biodiversidad de mayor tamaño (~US$300 M): confirmar cuál corresponde antes de citar cifras.')
bullet('No presentar EcuaLedger como criptomoneda ni activo especulativo: es infraestructura pública digital soberana con seguridad jurídica.')

doc.add_paragraph()
para('___________________________________', color=GREY, after=2)
para('Francisco Duque', bold=True, size=11, after=0)
para('Fundación Blockchain Soberana del Ecuador  ·  Quevedo, Ecuador  ·  Julio de 2026', size=9.5, color=GREY)

out = r'C:\Users\datos\Dropbox\var 91\Neg Inm\1 Proy Belen 2026\Tokenizacion Activos Reales - RWA\Fondos Disponibles Gracias\01_Borrador_Proyecto\BORRADOR_Proyecto_EcuaLedger_BID-VELA_v1.docx'
doc.save(out)
print('SAVED:', out)
